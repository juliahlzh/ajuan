import csv
import io
from datetime import datetime, timedelta
from typing import Any

from werkzeug.datastructures import FileStorage

from config import ALLOWED_IMPORT_EXTENSIONS

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


REQUIRED_COLUMNS = {
    "tanggal_ajuan",
    "nama_pengaju",
    "divisi_cabang",
    "jenis_ajuan",
    "nominal",
}

KEGIATAN_COLUMNS = {
    "Tanggal",
    "Divisi/Departemen",
    "Keterangan",
    "Ajuan",
    "Nominal",
    "Rekening an.",
}

SKIP_SHEETS = {
    "Rekap",
    "Lokasi",
    "Sheet4",
    "Sheet12",
    "Sheet2",
    "Promas_detail",
    "ART",
    "ART PUSAT",
}


def is_allowed_import_file(filename: str) -> bool:
    if "." not in filename:
        return False
    return filename.rsplit(".", 1)[1].lower() in ALLOWED_IMPORT_EXTENSIONS


def _to_iso_date(raw_value: Any) -> str:
    if isinstance(raw_value, datetime):
        if raw_value.year < 2020 or raw_value.year > 2030:
            return ""
        return raw_value.strftime("%Y-%m-%d")

    text = str(raw_value or "").strip()
    if not text:
        return ""

    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%m/%d/%Y",
        "%d.%m.%Y",
        "%Y/%m/%d",
    ]
    for fmt in formats:
        try:
            parsed = datetime.strptime(text, fmt)
            if parsed.year < 2020 or parsed.year > 2030:
                return ""
            return parsed.strftime("%Y-%m-%d")
        except ValueError:
            pass

    return ""


def _to_nominal(raw_value: Any) -> str:
    if raw_value is None:
        return "0"
    if isinstance(raw_value, (int, float)):
        return str(float(raw_value))
    text = str(raw_value).replace(",", "").strip()
    return text or "0"


def _calc_due_date(tanggal_ajuan: str, due_option: str, manual_due: str, wajib_lpj: str) -> str:
    if str(wajib_lpj).strip() not in {"1", "ya", "Ya", "YA", "true", "True", "wajib", "Wajib"}:
        return ""

    manual = _to_iso_date(manual_due)
    if manual:
        return manual

    if due_option not in {"30", "90", "365"}:
        due_option = "30"

    if not tanggal_ajuan:
        return ""

    base = datetime.strptime(tanggal_ajuan, "%Y-%m-%d")
    return (base + timedelta(days=int(due_option))).strftime("%Y-%m-%d")


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    tanggal_ajuan = _to_iso_date(row.get("tanggal_ajuan", ""))
    wajib_lpj = str(row.get("wajib_lpj", "1") or "1").strip()
    lpj_due_option = str(row.get("lpj_due_option", "30") or "30").strip()

    status_lpj_default = (
        "Belum Upload"
        if wajib_lpj in {"1", "ya", "Ya", "YA", "true", "True", "wajib", "Wajib"}
        else "Tidak Wajib"
    )

    return {
        "id_ajuan": str(row.get("id_ajuan", "") or "").strip(),
        "tanggal_ajuan": tanggal_ajuan,
        "nama_pengaju": str(row.get("nama_pengaju", "") or "").strip(),
        "nama_penerima_dana": str(
            row.get("nama_penerima_dana")
            or row.get("atas_nama_rekening")
            or row.get("rekening_an")
            or row.get("an_rekening_tujuan")
            or ""
        ).strip(),
        "nomor_memo": str(row.get("nomor_memo", "") or "").strip(),
        "perihal": str(row.get("perihal", "") or "").strip(),
        "divisi_cabang": str(row.get("divisi_cabang", "") or "").strip(),
        "jenis_ajuan": str(row.get("jenis_ajuan", "") or "").strip(),
        "nominal": _to_nominal(row.get("nominal", "0")),
        "keterangan": str(row.get("keterangan", "") or "").strip(),
        "status": str(row.get("status", "Menunggu Pembayaran") or "Menunggu Pembayaran").strip(),
        "nominal_realisasi": _to_nominal(row.get("nominal_realisasi", "")) if str(row.get("nominal_realisasi", "")).strip() else "",
        "link_bukti_transfer": str(row.get("link_bukti_transfer", "") or "").strip(),
        "link_memo_ajuan": str(row.get("link_memo_ajuan", "") or "").strip(),
        "wajib_lpj": wajib_lpj,
        "tanggal_jatuh_tempo_lpj": _calc_due_date(
            tanggal_ajuan=tanggal_ajuan,
            due_option=lpj_due_option,
            manual_due=str(row.get("tanggal_jatuh_tempo_lpj", "") or "").strip(),
            wajib_lpj=wajib_lpj,
        ),
        "status_lpj": str(row.get("status_lpj", status_lpj_default) or status_lpj_default).strip(),
        "link_lpj": str(row.get("link_lpj", "") or "").strip(),
    }


def _validate_headers(headers: list[str]) -> None:
    normalized = {h.strip() for h in headers if h and h.strip()}
    missing = REQUIRED_COLUMNS - normalized
    if missing:
        raise ValueError(f"Kolom wajib tidak ditemukan: {', '.join(sorted(missing))}")


def _is_kegiatan_layout(headers: list[str]) -> bool:
    normalized = {h.strip() for h in headers if h and h.strip()}
    return KEGIATAN_COLUMNS.issubset(normalized)


def _parse_kegiatan_sheet(rows_raw: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    last_tanggal = ""
    last_divisi = ""

    for raw in rows_raw:
        tanggal_ajuan = _to_iso_date(raw.get("Tanggal", ""))
        if tanggal_ajuan:
            last_tanggal = tanggal_ajuan
        else:
            tanggal_ajuan = last_tanggal

        divisi = str(raw.get("Divisi/Departemen", "") or "").strip()
        if divisi:
            last_divisi = divisi
        else:
            divisi = last_divisi

        keterangan = str(raw.get("Keterangan", "") or "").strip()
        nominal = _to_nominal(raw.get("Nominal", "0"))
        nama_penerima = str(raw.get("Rekening an.", "") or "").strip()
        memo_link = str(raw.get("Ajuan", "") or "").strip()
        wilayah = str(raw.get("Wilayah", "") or "").strip()

        if not any([tanggal_ajuan, divisi, keterangan, nominal, nama_penerima, memo_link]):
            continue
        if not tanggal_ajuan:
            continue

        rows.append(
            {
                "id_ajuan": "",
                "tanggal_ajuan": tanggal_ajuan,
                "nama_pengaju": "-",
                "nama_penerima_dana": nama_penerima or "",
                "nomor_memo": "",
                "perihal": keterangan,
                "divisi_cabang": divisi or wilayah or "-",
                "jenis_ajuan": "Kegiatan",
                "nominal": nominal,
                "keterangan": keterangan,
                "status": "Menunggu Pembayaran",
                "nominal_realisasi": "",
                "link_bukti_transfer": str(raw.get("Bukti TF", "") or "").strip(),
                "link_memo_ajuan": memo_link,
                "wajib_lpj": "1",
                "tanggal_jatuh_tempo_lpj": "",
                "status_lpj": "Belum Upload",
                "link_lpj": "",
            }
        )

    return rows


def _find_kpi_headers(ws) -> list[tuple[int, int]]:
    headers: list[tuple[int, int]] = []
    max_scan_row = min(ws.max_row, 80)
    max_scan_col = min(ws.max_column, 80)

    for r in range(1, max_scan_row + 1):
        for c in range(1, max_scan_col + 1):
            a = str(ws.cell(r, c).value or "").strip().lower()
            b = str(ws.cell(r, c + 1).value or "").strip().lower()
            d = str(ws.cell(r, c + 2).value or "").strip().lower()
            if a == "tanggal ajuan" and b == "keterangan" and d == "ajuan":
                headers.append((r, c))
    return headers


def _parse_kpi_table(ws, header_row: int, start_col: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    empty_streak = 0
    last_tanggal = ""

    for r in range(header_row + 1, ws.max_row + 1):
        tanggal_raw = ws.cell(r, start_col).value
        ket_raw = ws.cell(r, start_col + 1).value
        ajuan_raw = ws.cell(r, start_col + 2).value
        tgl_tf_raw = ws.cell(r, start_col + 3).value
        nama_raw = ws.cell(r, start_col + 4).value
        realisasi_raw = ws.cell(r, start_col + 5).value
        tgl_lpj_raw = ws.cell(r, start_col + 6).value
        wilayah_raw = ws.cell(r, start_col + 7).value

        if not any(v not in (None, "") for v in [tanggal_raw, ket_raw, ajuan_raw, nama_raw, wilayah_raw]):
            empty_streak += 1
            if empty_streak >= 25:
                break
            continue

        empty_streak = 0

        tanggal_ajuan = _to_iso_date(tanggal_raw)
        if tanggal_ajuan:
            last_tanggal = tanggal_ajuan
        else:
            tanggal_ajuan = last_tanggal

        keterangan = str(ket_raw or "").strip()
        nominal = _to_nominal(ajuan_raw)
        nama_penerima = str(nama_raw or "").strip() or ""
        wilayah = str(wilayah_raw or "").strip()

        if not tanggal_ajuan and not keterangan:
            continue

        rows.append(
            {
                "id_ajuan": "",
                "tanggal_ajuan": tanggal_ajuan,
                "nama_pengaju": "-",
                "nama_penerima_dana": nama_penerima,
                "nomor_memo": "",
                "perihal": keterangan,
                "divisi_cabang": wilayah or ws.title,
                "jenis_ajuan": ws.title,
                "nominal": nominal,
                "keterangan": keterangan,
                "status": "Selesai Dibayar" if _to_iso_date(tgl_tf_raw) else "Menunggu Pembayaran",
                "nominal_realisasi": _to_nominal(realisasi_raw) if str(realisasi_raw or "").strip() else "",
                "link_bukti_transfer": "",
                "link_memo_ajuan": "",
                "wajib_lpj": "1",
                "tanggal_jatuh_tempo_lpj": "",
                # Tetap Belum Upload agar LPJ 2025 bisa diinput di aplikasi.
                "status_lpj": "Belum Upload",
                "link_lpj": "",
                "tanggal_lpj_sumber": _to_iso_date(tgl_lpj_raw),
            }
        )

    return rows


def _parse_kpi_workbook(wb) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen_keys: set[tuple[str, str, str, str, str]] = set()

    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue

        header_positions = _find_kpi_headers(ws)
        if not header_positions:
            continue

        for header_row, start_col in header_positions:
            parsed = _parse_kpi_table(ws, header_row, start_col)
            for item in parsed:
                key = (
                    item.get("tanggal_ajuan", ""),
                    item.get("keterangan", ""),
                    item.get("nominal", ""),
                    item.get("divisi_cabang", ""),
                    item.get("nama_pengaju", ""),
                )
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                rows.append(item)

    return rows


def parse_import_file(file_obj: FileStorage) -> list[dict[str, Any]]:
    filename = file_obj.filename or ""
    ext = filename.rsplit(".", 1)[1].lower()

    if ext == "csv":
        raw = file_obj.stream.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(raw))
        headers = reader.fieldnames or []
        _validate_headers(headers)
        rows = [_normalize_row(r) for r in reader if any((v or "").strip() for v in r.values())]
        return rows

    if ext == "xlsx":
        if load_workbook is None:
            raise ValueError("File XLSX butuh dependency openpyxl. Jalankan: pip install openpyxl")

        wb = load_workbook(file_obj.stream, data_only=True)

        if "Kegiatan" in wb.sheetnames:
            ws = wb["Kegiatan"]
            headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            if _is_kegiatan_layout(headers):
                rows_raw: list[dict[str, Any]] = []
                for idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
                    if idx == 1:
                        continue

                    row_map = {}
                    empty = True
                    for i, header in enumerate(headers):
                        if not header:
                            continue
                        cell = values[i] if i < len(values) else ""
                        if isinstance(cell, datetime):
                            cell = cell.strftime("%Y-%m-%d")
                        elif isinstance(cell, (int, float)) and header in {"tanggal_ajuan", "Tanggal"}:
                            base = datetime(1899, 12, 30)
                            cell = (base + timedelta(days=float(cell))).strftime("%Y-%m-%d")
                        else:
                            cell = "" if cell is None else str(cell)

                        if str(cell).strip():
                            empty = False
                        row_map[header] = cell

                    if not empty:
                        rows_raw.append(row_map)

                return _parse_kegiatan_sheet(rows_raw)

        kpi_rows = _parse_kpi_workbook(wb)
        if kpi_rows:
            return kpi_rows

        ws = wb.active
        headers: list[str] = []
        rows: list[dict[str, Any]] = []
        for idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx == 1:
                headers = [str(v).strip() if v is not None else "" for v in values]
                _validate_headers(headers)
                continue

            row_map = {}
            empty = True
            for i, header in enumerate(headers):
                if not header:
                    continue
                cell = values[i] if i < len(values) else ""
                if isinstance(cell, datetime):
                    cell = cell.strftime("%Y-%m-%d")
                elif isinstance(cell, (int, float)) and header == "tanggal_ajuan":
                    base = datetime(1899, 12, 30)
                    cell = (base + timedelta(days=float(cell))).strftime("%Y-%m-%d")
                else:
                    cell = "" if cell is None else str(cell)

                if str(cell).strip():
                    empty = False
                row_map[header] = cell

            if not empty:
                rows.append(_normalize_row(row_map))

        return rows

    raise ValueError("Format import tidak didukung. Gunakan CSV atau XLSX.")
